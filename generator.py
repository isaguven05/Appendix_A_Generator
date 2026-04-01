"""
Core PPTX generation logic for Appendix A automation.
"""

import io
import os
import re
import uuid
import zipfile
import html as html_module
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor


# ─── Shape bounds from template (in EMUs) ─────────────────────────────────────
# Slide 1 – Title: image rectangle
TITLE_IMAGE_BOUNDS = dict(left=349371, top=3255351, width=14414258, height=6462572)
# RF Schematics: image rectangle
RF_IMAGE_BOUNDS    = dict(left=262515,  top=2511339, width=11871901, height=6783229)
# GA: image rectangle
GA_IMAGE_BOUNDS    = dict(left=346087,  top=1890006, width=11818354, height=6989657)
# GA: title text box
GA_TITLE_BOUNDS    = dict(left=435937,  top=616155,  width=3598287,  height=634118)


# ─── Helpers ──────────────────────────────────────────────────────────────────

def xml_escape(text: str) -> str:
    return (str(text)
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace('"', "&quot;"))


def extract_station_info(image_3d_path: str):
    """Parse '3D_DO93_TEMPLE.png' → (code='DO93', name='TEMPLE', combined='DO93_TEMPLE')."""
    stem = Path(image_3d_path).stem          # '3D_DO93_TEMPLE'
    parts = stem.split("_", 2)               # ['3D', 'DO93', 'TEMPLE']
    if len(parts) < 3:
        raise ValueError(f"Unexpected 3D image filename: {image_3d_path!r}")
    code = parts[1]
    name = parts[2]
    combined = f"{code}_{name}"
    return code, name, combined


def get_initials(full_name: str) -> str:
    """'Aleem Alvi' → 'AA'"""
    return "".join(p[0].upper() for p in full_name.strip().split() if p)


def sort_design_plans(paths: list) -> list:
    """Sort design-plan images by trailing number in filename."""
    def _key(p):
        m = re.search(r"(\d+)[^0-9]*$", Path(p).stem)
        return int(m.group(1)) if m else 0
    return sorted(paths, key=_key)


def extract_building_title(filename: str) -> str:
    """'EE_TEMPLE - Building 1 - Midway.jpeg' → 'Midway'
    Also handles secure_filename output: 'EE_TEMPLE_-_Building_1_-_Midway.jpeg'"""
    stem = Path(filename).stem
    stem_normalised = stem.replace("_-_", " - ")
    parts = stem_normalised.split(" - ")
    title = parts[-1].strip() if len(parts) >= 2 else stem
    title = title.replace("_", " ")
    return title


# ─── python-pptx slide helpers ────────────────────────────────────────────────

def delete_slide(prs: Presentation, index: int):
    """Delete slide at index from a Presentation (in-place)."""
    xml_slides = prs.slides._sldIdLst
    elem = xml_slides[index]
    r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    rId = elem.get(f"{{{r_ns}}}id")
    prs.part.drop_rel(rId)
    xml_slides.remove(elem)


def add_image_shape(slide, bounds: dict, image_path: str):
    """Add a picture to slide at given EMU bounds dict (stretches to fill)."""
    slide.shapes.add_picture(
        image_path,
        Emu(bounds["left"]),
        Emu(bounds["top"]),
        Emu(bounds["width"]),
        Emu(bounds["height"]),
    )


def _fit_image(img_w, img_h, box_w, box_h):
    """
    Scale (img_w × img_h) to fit inside (box_w × box_h) preserving aspect ratio.
    Returns (scaled_w, scaled_h, left_offset, top_offset) — all in EMUs — so
    the caller can centre the result inside the box.
    """
    scale = min(box_w / img_w, box_h / img_h)
    w = int(img_w * scale)
    h = int(img_h * scale)
    dx = (box_w - w) // 2
    dy = (box_h - h) // 2
    return w, h, dx, dy


def add_image_fitted(slide, bounds: dict, image_path: str):
    """
    Add a picture fitted (aspect-ratio-preserved, centred) inside bounds dict.
    Uses the same approach as Appendix B's _place_image_fitted / update_slide.
    """
    box_left = Emu(bounds["left"])
    box_top  = Emu(bounds["top"])
    box_w    = Emu(bounds["width"])
    box_h    = Emu(bounds["height"])

    # Add temporarily at 0,0 to get python-pptx's natural EMU dimensions
    tmp = slide.shapes.add_picture(image_path, 0, 0)
    nat_w, nat_h = tmp.width, tmp.height
    slide.shapes._spTree.remove(tmp.element)

    w, h, dx, dy = _fit_image(nat_w, nat_h, box_w, box_h)
    slide.shapes.add_picture(image_path, box_left + dx, box_top + dy,
                             width=w, height=h)


def add_text_box(slide, bounds: dict, text: str, font_size_pt: int = 20,
                 bold: bool = True):
    """Add a plain text box to a slide at given EMU bounds dict."""
    txBox = slide.shapes.add_textbox(
        Emu(bounds["left"]),
        Emu(bounds["top"]),
        Emu(bounds["width"]),
        Emu(bounds["height"]),
    )
    tf = txBox.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    return txBox


# ─── XML-level post-processing ────────────────────────────────────────────────

def _replace_in_bytes(content: bytes, replacements: dict) -> bytes:
    """
    Replace XML-escaped placeholder strings in raw XML bytes.
    Both the escaped form (&lt;&lt;FOO&gt;&gt;) and unescaped form (<<FOO>>)
    are searched; replacement values are XML-escaped automatically.
    """
    text = content.decode("utf-8")
    for raw_key, raw_value in replacements.items():
        if raw_key is None or raw_value is None:
            continue
        escaped_key = xml_escape(raw_key)
        escaped_value = xml_escape(str(raw_value))
        text = text.replace(escaped_key, escaped_value)
        if raw_key != escaped_key:
            text = text.replace(raw_key, escaped_value)
    return text.encode("utf-8")


def _install_slidenum_field(content: bytes) -> bytes:
    """
    Replace literal <NUM> text run in master/layout XML with a PPTX slidenum field.
    """
    text = content.decode("utf-8")
    field_id = "{" + str(uuid.uuid4()).upper() + "}"
    text = re.sub(
        r"<a:r>(<a:rPr[^/]*/>\s*)?<a:t>&lt;NUM&gt;</a:t></a:r>",
        f'<a:fld id="{field_id}" type="slidenum">'
        f'<a:rPr lang="en-GB" smtClean="0" dirty="0"/>'
        f"<a:t>\u2039#\u203a</a:t>"
        f"</a:fld>",
        text,
    )
    return text.encode("utf-8")


def _post_process_pptx(pptx_bytes: bytes, replacements: dict) -> bytes:
    """
    Open PPTX bytes as zip, apply text replacements across master, layouts,
    and all slides, then return modified bytes.
    """
    in_buf = io.BytesIO(pptx_bytes)
    out_buf = io.BytesIO()

    with zipfile.ZipFile(in_buf, "r") as zin, \
         zipfile.ZipFile(out_buf, "w", zipfile.ZIP_DEFLATED) as zout:

        for item in zin.infolist():
            data = zin.read(item.filename)

            if item.filename.endswith(".xml"):
                is_master = "slideMasters/" in item.filename
                is_layout = "slideLayouts/" in item.filename
                is_slide  = re.match(r".*ppt/slides/slide\d+\.xml$", item.filename)

                if is_master or is_layout or is_slide:
                    data = _replace_in_bytes(data, replacements)

                if is_master or is_layout:
                    data = _install_slidenum_field(data)

            zout.writestr(item, data)

    return out_buf.getvalue()


# ─── Main generation function ─────────────────────────────────────────────────

def generate_pptx(
    template_path: str,
    image_3d_path: str,
    p01_date: str,
    p02_date: str,
    author: str,
    checked: str,
    approved: str,
    station_address: str,
    design_plan_images: list,      # file paths sorted by plan number
    building_images: list,         # list of file paths
    building_titles: list = None,  # optional pre-extracted titles (preserves '&' etc.)
) -> bytes:
    """
    Generate the Appendix A PPTX for one provider and return it as bytes.

    Output slide order:
      1  – Title slide
      2  – Design Change Summary
      3  – 2_GA slide  (always exactly 2)
      4  – 2_GA slide
      5… – RF Schematics slides (one per design plan image)
      …  – GA slides (one per building image)
    """

    # ── Station metadata ──────────────────────────────────────────────────────
    code, station_name, combined = extract_station_info(image_3d_path)
    d_ini = get_initials(author)
    c_ini = get_initials(checked)
    a_ini = get_initials(approved)

    design_plan_images = sort_design_plans(design_plan_images)

    # ── Load template ─────────────────────────────────────────────────────────
    prs = Presentation(template_path)

    # Build layout map – use first occurrence so duplicate names don't overwrite
    layout_map = {}
    for lay in prs.slide_layouts:
        if lay.name not in layout_map:
            layout_map[lay.name] = lay

    lay_2ga = layout_map["2_GA"]
    lay_rf  = layout_map["RF Schematics"]
    lay_ga  = layout_map["GA"]

    # Template has 5 slides; keep slides 0 & 1 (Title + Design Change Summary),
    # delete the rest
    while len(prs.slides) > 2:
        delete_slide(prs, 2)

    # ── Always exactly 2 × 2_GA slides ───────────────────────────────────────
    for _ in range(2):
        prs.slides.add_slide(lay_2ga)

    # ── RF Schematics slides ──────────────────────────────────────────────────
    rf_slides = []
    for img_path in design_plan_images:
        slide = prs.slides.add_slide(lay_rf)
        rf_slides.append((slide, img_path))

    # ── GA (building) slides ──────────────────────────────────────────────────
    ga_slides = []
    for img_path in building_images:
        slide = prs.slides.add_slide(lay_ga)
        ga_slides.append((slide, img_path))

    # ── Title slide – replace 3D image placeholder ────────────────────────────
    title_slide = prs.slides[0]
    for shape in list(title_slide.shapes):
        # Remove the placeholder text box AND any background rectangle/image
        # so the 3D image sits cleanly without anything behind it
        if shape.has_text_frame and "<<3D_IMAGE>>" in shape.text_frame.text:
            shape._element.getparent().remove(shape._element)
        elif shape.shape_type in (1, 13):  # 1=rectangle, 13=picture (background shapes)
            left, top = shape.left, shape.top
            bnd_left = TITLE_IMAGE_BOUNDS["left"]
            bnd_top  = TITLE_IMAGE_BOUNDS["top"]
            bnd_right  = bnd_left + TITLE_IMAGE_BOUNDS["width"]
            bnd_bottom = bnd_top  + TITLE_IMAGE_BOUNDS["height"]
            # Only remove shapes that sit inside or overlap the image area
            if left >= bnd_left - 200000 and top >= bnd_top - 200000 \
                    and left <= bnd_right and top <= bnd_bottom:
                shape._element.getparent().remove(shape._element)
    add_image_fitted(title_slide, TITLE_IMAGE_BOUNDS, image_3d_path)

    # ── RF Schematics slides – insert cable images ────────────────────────────
    for slide, img_path in rf_slides:
        add_image_fitted(slide, RF_IMAGE_BOUNDS, img_path)

    # ── GA slides – insert building images and titles ─────────────────────────
    for idx, (slide, img_path) in enumerate(ga_slides):
        if building_titles and idx < len(building_titles):
            bld_title = building_titles[idx]
        else:
            bld_title = extract_building_title(img_path)
        add_image_fitted(slide, GA_IMAGE_BOUNDS, img_path)
        add_text_box(slide, GA_TITLE_BOUNDS, bld_title, font_size_pt=24, bold=True)

    # ── Save to bytes ─────────────────────────────────────────────────────────
    buf = io.BytesIO()
    prs.save(buf)
    pptx_bytes = buf.getvalue()

    # ── XML post-processing: text replacements + slidenum field ───────────────
    replacements = {
        "<<STATION_NAME>>": combined,
        "<<ADDRESS>>":      station_address,
        "<P01>":            p01_date,
        "<P02>":            p02_date,
        "<D>":              d_ini,
        "<C>":              c_ini,
        "<A>":              a_ini,
        "<<STATION>>":      f"{code} {station_name}",
        "<<CODE>>":         code,
        "<<Author>>":       author,
        "<<Checked>>":      checked,
        "<<Approved>>":     approved,
    }

    pptx_bytes = _post_process_pptx(pptx_bytes, replacements)
    return pptx_bytes
