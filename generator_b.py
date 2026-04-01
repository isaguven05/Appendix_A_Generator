"""
Core PPTX generation logic for Appendix B.

Uses the building logic from engine_b.py (the original engine.py) without
any modification.  The extraction phase (reading from an input PPTX) is
handled in app.py via engine_b.process_pptx.

Provider codes for the <P> placeholder:
    EE → 2   VODAFONE → 3   VMO2 → 4   THREE → 5
"""

import io
import os
import re
import shutil
import tempfile
import zipfile as _zf

from pptx import Presentation

# ── Import building helpers from engine_b (unmodified engine.py) ──────────────
from engine_b import (
    TEMPLATE_MAP,
    DEFAULT_BUILDING_ORDER,
    _make_sort_key,
    duplicate_slide,
    update_slide,
    add_intro_slides,
    delete_template_slides,
    replace_in_master,
    replace_logo_in_master,
    _initials,
)

# ── Import XML helpers from generator.py ─────────────────────────────────────
from generator import extract_station_info, _replace_in_bytes

# Provider code mapping — inserted into <P> placeholder on every slide
PROVIDER_CODES = {
    "EE":       "2",
    "VODAFONE": "3",
    "VMO2":     "4",
    "THREE":    "5",
}


# ─── Custom post-processor (no slidenum field injection) ──────────────────────
# templateB already has proper <a:fld type="slidenum"/> in every layout.
# _post_process_pptx from generator.py also calls _install_slidenum_field on
# masters/layouts, which adds a SECOND field — causing numbers to show doubled
# (e.g. slide 1 renders as "11").  This version skips that step.

def _strip_num_placeholder(data: bytes) -> bytes:
    """
    Remove any residual <NUM> / &lt;NUM&gt; text runs from slide/layout/master XML.
    The template already has proper <a:fld type="slidenum"/> fields everywhere;
    these stray <NUM> runs would otherwise leave blank space or confuse renderers.
    """
    text = data.decode("utf-8")
    # Remove the whole <a:r>…<a:t>&lt;NUM&gt;</a:t>…</a:r> run (rPr is optional)
    text = re.sub(
        r"<a:r>(?:<a:rPr[^>]*(?:/>|>.*?</a:rPr>))?\s*<a:t>&lt;NUM&gt;</a:t></a:r>",
        "",
        text,
        flags=re.DOTALL,
    )
    # Safety: also remove any bare &lt;NUM&gt; that survived the above
    text = text.replace("&lt;NUM&gt;", "")
    return text.encode("utf-8")


def _b_post_process_pptx(pptx_bytes: bytes, replacements: dict) -> bytes:
    """
    Apply text replacements to all slides, masters and layouts in pptx_bytes.
    Also strips residual <NUM> placeholder runs.
    Does NOT inject new slide-number fields (templateB already has them).
    """
    in_buf  = io.BytesIO(pptx_bytes)
    out_buf = io.BytesIO()

    with _zf.ZipFile(in_buf, "r") as zin, \
         _zf.ZipFile(out_buf, "w", _zf.ZIP_DEFLATED) as zout:

        for item in zin.infolist():
            data = zin.read(item.filename)

            if item.filename.endswith(".xml"):
                is_master = "slideMasters/" in item.filename
                is_layout = "slideLayouts/" in item.filename
                is_slide  = bool(
                    re.match(r".*ppt/slides/slide\d+\.xml$", item.filename)
                )
                if is_master or is_layout or is_slide:
                    data = _replace_in_bytes(data, replacements)
                    data = _strip_num_placeholder(data)   # ← remove <NUM> runs

            zout.writestr(item, data)

    return out_buf.getvalue()


# ─── ZIP-level master logo injection ─────────────────────────────────────────
# python-pptx's MasterShapes does not expose add_picture(), so we manipulate
# the PPTX zip directly after saving.

def _add_logo_to_master_zip(pptx_bytes: bytes, logo_path: str) -> bytes:
    """
    Pure string-manipulation approach: inject the provider logo into every
    slide in the PPTX so it appears top-right on all slides.

    Uses raw string operations (no lxml re-serialisation) to avoid any XML
    namespace or encoding changes that could cause PowerPoint to reject the file.

    Fixed position/size (pt → EMU, 1 pt = 12700 EMU):
      X = 1050 pt → 13 335 000 EMU
      Y =   62 pt →    787 400 EMU
      W =   43 pt →    546 100 EMU
      H =   43 pt →    546 100 EMU
    """
    import re as _re

    has_logo = bool(logo_path and os.path.exists(logo_path))
    if not has_logo:
        return pptx_bytes

    with open(logo_path, 'rb') as fh:
        logo_bytes = fh.read()
    logo_ext   = os.path.splitext(logo_path)[1].lower()
    mime       = {'.png': 'image/png', '.jpg': 'image/jpeg',
                  '.jpeg': 'image/jpeg', '.gif': 'image/gif'}.get(logo_ext, 'image/png')
    media_zip  = f'ppt/media/prov_logo_slide{logo_ext}'
    # Relative path from a slide's _rels file to the media folder
    rel_target = f'../media/prov_logo_slide{logo_ext}'
    IMG_REL    = ('http://schemas.openxmlformats.org/officeDocument/'
                  '2006/relationships/image')
    RID        = 'rId_prov_logo'

    L, T, W, H = 13_335_000, 787_400, 546_100, 546_100

    # The <p:pic> XML to inject — uses the same namespace prefixes already
    # declared in every slide XML so no extra xmlns declarations are needed.
    PIC = (
        f'<p:pic>'
          f'<p:nvPicPr>'
            f'<p:cNvPr id="9997" name="ProviderLogo"/>'
            f'<p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr>'
            f'<p:nvPr/>'
          f'</p:nvPicPr>'
          f'<p:blipFill>'
            f'<a:blip r:embed="{RID}"/>'
            f'<a:stretch><a:fillRect/></a:stretch>'
          f'</p:blipFill>'
          f'<p:spPr>'
            f'<a:xfrm><a:off x="{L}" y="{T}"/><a:ext cx="{W}" cy="{H}"/></a:xfrm>'
            f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
          f'</p:spPr>'
        f'</p:pic>'
    )
    REL_ENTRY = (
        f'<Relationship Id="{RID}" Type="{IMG_REL}" Target="{rel_target}"/>'
    )

    in_buf  = io.BytesIO(pptx_bytes)
    out_buf = io.BytesIO()

    with _zf.ZipFile(in_buf, 'r') as zin, \
         _zf.ZipFile(out_buf, 'w', _zf.ZIP_DEFLATED) as zout:

        for item in zin.infolist():
            data = zin.read(item.filename)
            fn   = item.filename

            # ── Individual slides: inject <p:pic> into spTree ────────────────
            if _re.match(r'ppt/slides/slide\d+\.xml$', fn):
                text = data.decode('utf-8')
                if RID not in text:          # don't double-inject
                    # Insert just before the closing </p:spTree>
                    text = text.replace('</p:spTree>', PIC + '</p:spTree>', 1)
                data = text.encode('utf-8')

            # ── Slide rels: add image relationship ───────────────────────────
            elif _re.match(r'ppt/slides/_rels/slide\d+\.xml\.rels$', fn):
                text = data.decode('utf-8')
                if RID not in text:
                    text = text.replace('</Relationships>',
                                        REL_ENTRY + '</Relationships>')
                data = text.encode('utf-8')

            # ── Content types: register extension if new ─────────────────────
            elif fn == '[Content_Types].xml':
                text     = data.decode('utf-8')
                ext_bare = logo_ext.lstrip('.')
                if f'Extension="{ext_bare}"' not in text:
                    entry = (f'<Default Extension="{ext_bare}" '
                             f'ContentType="{mime}"/>')
                    text = text.replace('</Types>', entry + '</Types>')
                data = text.encode('utf-8')

            zout.writestr(item, data)

        # Add the logo image once into the zip
        zout.writestr(media_zip, logo_bytes)

    return out_buf.getvalue()


# ─── Main generation function ─────────────────────────────────────────────────

def generate_pptx_b(
    template_path: str,
    image_3d_path: str,
    date: str,
    p01_date: str,
    author: str,
    checked: str,
    approved: str,
    address: str,
    image_entries: list,      # [(saved_path, original_name), …] for ONE provider
    provider: str,
    building_order: list = None,
    logo_path: str = None,
) -> bytes:
    """
    Generate one Appendix B PPTX for a single provider and return as bytes.

    Slide structure (driven by templateB.pptx + engine_b):
      Slide 1  – Title slide (3D station image)
      Slide 2  – Executive summary
      Slide 3  – Operator logo slide
      Slides 4+ – RF prediction maps, sorted by tech/freq/metric + building

    Metadata replacements:
      Via engine_b replace_in_master:  <<d>> <<c>> <<a>> <<po1>> <<po2>>
      Via _b_post_process_pptx:        <<DATE>> <<AUTHOR>> <<CHECKED>>
                                        <<APPROVED>> <<ADDRESS>> <<P01>>
                                        <<P02>> <<STATION_NAME>> <D> <C>
                                        <A> <P>  STATION-NAME  OPERATOR-NAME
    """

    # ── Station info from 3D image filename ───────────────────────────────────
    try:
        code, name, combined = extract_station_info(image_3d_path)
        # name may contain underscores (e.g. "LAMBETH_NORTH")
        name_spaced   = name.replace("_", " ").title()      # "Lambeth North"
        station_label = f"{code} {name_spaced}"             # "B137 Lambeth North"
    except Exception:
        code          = "UNKNOWN"
        name          = "UNKNOWN"
        combined      = "UNKNOWN"
        name_spaced   = "UNKNOWN"
        station_label = "UNKNOWN"

    # ── Metadata dict used by engine_b helpers ────────────────────────────────
    metadata = {
        "date":     date,
        "author":   author,
        "checked":  checked,
        "approved": approved,
        "p01_date": p01_date,
    }

    # ── Parse image entries into provider_slides list ─────────────────────────
    # Filename format: PROVIDER_TECH_MHZ_METRIC_BUILDING...png
    provider_slides = []

    for saved_path, orig_name in image_entries:
        stem    = os.path.splitext(orig_name)[0]
        parts_f = stem.replace(" ", "_").split("_")

        if len(parts_f) < 5:
            continue

        tech     = parts_f[1].upper()
        mhz      = parts_f[2]
        metric   = parts_f[3].upper()
        building = "_".join(parts_f[4:])
        key      = (tech, mhz, metric)

        if key not in TEMPLATE_MAP:
            continue

        provider_slides.append(
            (provider, orig_name, str(saved_path), tech, mhz, metric, building)
        )

    # ── Sort by tech/freq/metric then building order ──────────────────────────
    sort_key = _make_sort_key(building_order or DEFAULT_BUILDING_ORDER)
    provider_slides.sort(key=sort_key)

    # ── Build the presentation in a temp dir ──────────────────────────────────
    with tempfile.TemporaryDirectory() as tmpdir:
        out_path = os.path.join(tmpdir, "output.pptx")

        shutil.copy(template_path, out_path)
        prs  = Presentation(out_path)
        tmpl = Presentation(template_path)

        n_template_slides = len(prs.slides)

        # ── Rename 3D image so add_intro_slides parses it correctly ───────────
        # add_intro_slides expects:  CODE_NAME_3DMOD.png
        #   parts[0]     → station code   (e.g. "B137")
        #   parts[1:-1]  → station name   (e.g. ["LAMBETH", "NORTH"])
        # Our upload format is  3D_CODE_NAME.png, so parts[0] = "3D" (wrong).
        # Copy to a temp file with the correct naming before passing it in.
        threed_name_fixed = f"{code}_{name}_3DMOD.png"
        threed_path_fixed = os.path.join(tmpdir, threed_name_fixed)
        shutil.copy(image_3d_path, threed_path_fixed)

        # ── Intro slides: title, exec summary, operator logo ──────────────────
        add_intro_slides(
            prs, tmpl,
            threed_image_path = threed_path_fixed,
            provider          = provider,
            logo_path         = logo_path or "",
            metadata          = metadata,
        )

        # ── Fix: place 3D image from GROUP placeholder in title slide ─────────
        # The template's title slide wraps the image area in a GROUP shape.
        # add_intro_slides only iterates top-level shapes and searches for
        # "3D MODEL IMAGE" text — it can't see inside the GROUP, so the image
        # is never placed.  We handle it here instead.
        slides_list = list(prs.slides)
        title_slide = slides_list[n_template_slides]   # first slide added

        _place_3d_image(title_slide, image_3d_path)

        # ── Content slides: one per RF map image ──────────────────────────────
        for _, _fname, path, tech, mhz, metric, building in provider_slides:
            tmpl_slide = tmpl.slides[TEMPLATE_MAP[(tech, mhz, metric)]]
            slide      = duplicate_slide(prs, tmpl_slide)
            update_slide(slide, path, building, logo_path or "")

        # ── Remove original blank template slides ─────────────────────────────
        delete_template_slides(prs, n_template_slides)

        # ── Logo: big centred logo on operator slide (slide index 2) ────────
        if logo_path and os.path.exists(logo_path):
            from engine_b import replace_logo_placeholder
            logo_slide = list(prs.slides)[2]
            replace_logo_placeholder(logo_slide, logo_path, big=True)

        # ── Master metadata (engine_b placeholders) ───────────────────────────
        replace_in_master(prs, {
            "<<d>>":   _initials(author),
            "<<c>>":   _initials(checked),
            "<<a>>":   _initials(approved),
            "<<po1>>": p01_date,
            "<<po2>>": date,
        })

        # ── Master logo: replace <<logo>> on slide master so it appears on
        #    every slide at exactly 1050pt x 62pt, 43×43pt ─────────────────
        replace_logo_in_master(prs, logo_path or "")

        # ── Save to bytes ─────────────────────────────────────────────────────
        buf = io.BytesIO()
        prs.save(buf)
        pptx_bytes = buf.getvalue()

    # ── XML post-processing: remaining placeholder replacements ───────────────
    d_ini = _initials(author)
    c_ini = _initials(checked)
    a_ini = _initials(approved)

    replacements = {
        "<<DATE>>":         date,
        "<<AUTHOR>>":       author,
        "<<CHECKED>>":      checked,
        "<<APPROVED>>":     approved,
        "<<ADDRESS>>":      address,
        "<<P01>>":          p01_date,
        "<<P02>>":          date,
        "<P01>":            p01_date,   # template uses single-bracket form
        "<P02>":            date,
        # Clean human-readable station name: "B137 Lambeth North" (no underscores)
        "<<STATION_NAME>>": station_label,
        "<D>":              d_ini,
        "<C>":              c_ini,
        "<A>":              a_ini,
        "<<d>>":            d_ini,
        "<<c>>":            c_ini,
        "<<a>>":            a_ini,
        "<<po1>>":          p01_date,
        "<<po2>>":          date,
        "<P>":              PROVIDER_CODES.get(provider.upper(), ""),
        "STATION-NAME":     station_label,
        "OPERATOR-NAME":    provider,
    }

    pptx_bytes = _b_post_process_pptx(pptx_bytes, replacements)
    pptx_bytes = _add_logo_to_master_zip(pptx_bytes, logo_path or "")
    return pptx_bytes


def _place_3d_image(slide, image_path: str):
    """
    Find the GROUP shape on the title slide that contains the 3D image area
    and replace it with the actual 3D station image.

    The template wraps the image placeholder in a GROUP (named 'Rectangle 5').
    add_intro_slides can't find it because it only iterates top-level shapes.
    This function removes the group and inserts the image at the same bounds.
    If no GROUP shape is found (or no image path given), nothing happens.
    """
    if not image_path or not os.path.exists(image_path):
        return

    for shape in list(slide.shapes):
        if shape.shape_type == 6:   # 6 = GROUP
            left, top   = shape.left, shape.top
            width, height = shape.width, shape.height
            slide.shapes._spTree.remove(shape.element)
            slide.shapes.add_picture(image_path, left, top, width=width, height=height)
            return   # one GROUP per title slide — done
